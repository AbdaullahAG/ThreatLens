"""
Integration tests for reporters and the SQLite enrichment cache.

Reporter formula-neutralisation
────────────────────────────────
Verifies that cell values beginning with spreadsheet formula prefixes
(=, +, -, @, tab, CR, LF) are prefixed with an apostrophe so they are
rendered as plain text — never executed — in Excel or any CSV reader.

SQLite cache integration
────────────────────────
Verifies that EnrichmentResult objects survive a full to_dict / from_dict
round-trip through InvestigationStore for each IOC type (IP, CVE).
"""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.models import IOC, IOCType, EnrichmentResult
from src.reporters.excel_reporter import ExcelReporter
from src.reporters.other_reporters import CSVReporter
from src.storage import InvestigationStore


# ── Helpers ───────────────────────────────────────────────────────────────────

def _cve_result(value: str = "CVE-2021-44228") -> EnrichmentResult:
    """Build a fully-populated CVE EnrichmentResult."""
    r = EnrichmentResult(ioc=IOC(value=value, ioc_type=IOCType.CVE))
    r.cvss_score = 10.0
    r.severity = "CRITICAL"
    r.cve_description = "Log4Shell remote code execution"
    r.published_date = "2021-12-10"
    r.sources["nvd"] = {"cvss_score": 10.0}
    r.set_verdict()
    return r


def _formula_result(formula: str) -> EnrichmentResult:
    """Build a result whose IOC value starts with a formula prefix."""
    r = EnrichmentResult(ioc=IOC(value=formula, ioc_type=IOCType.DOMAIN))
    r.malicious_votes = 5
    r.sources["test"] = {"malicious": 5}
    r.set_verdict()
    return r


# ── Excel formula neutralisation ─────────────────────────────────────────────

class TestExcelFormulaProtection:

    @pytest.mark.parametrize("formula", [
        "=HYPERLINK(\"https://evil.example\")",
        "+SUM(A1:A10)",
        "-1+2*SUM(A1)",
        "@SUM(A1:A10)",
        "\t=cmd",
        "\r=cmd",
        "\n=cmd",
    ])
    def test_formula_cell_is_neutralised(self, tmp_path: Path, formula: str):
        """Formula-starting strings must be apostrophe-prefixed in the workbook."""
        results = [_formula_result(formula)]
        filepath = ExcelReporter().generate(results, output_dir=str(tmp_path))

        wb = load_workbook(filepath)
        # Check all sheets for any cell that contains the raw formula (unexpected)
        found_raw = False
        found_safe = False
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value == formula:
                        found_raw = True
                    if isinstance(cell.value, str) and cell.value.startswith("'"):
                        found_safe = True

        assert not found_raw, f"Raw formula '{formula}' was written unprotected to Excel"
        assert found_safe, f"No apostrophe-prefixed cell found for formula '{formula}'"

    def test_numeric_values_are_not_prefixed(self, tmp_path: Path):
        """Plain integers must not be converted to strings."""
        r = _cve_result()
        filepath = ExcelReporter().generate([r], output_dir=str(tmp_path))
        wb = load_workbook(filepath)
        numeric_found = any(
            isinstance(cell.value, (int, float))
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
        )
        assert numeric_found, "No numeric cell found — spreadsheet_value() is over-converting"


# ── CSV formula neutralisation ────────────────────────────────────────────────

class TestCSVFormulaProtection:

    @pytest.mark.parametrize("formula", [
        "=CMD",
        "+1+1",
        "-1*2",
        "@SUM(A1)",
    ])
    def test_formula_cell_is_neutralised(self, tmp_path: Path, formula: str):
        """Formula-starting strings must be apostrophe-prefixed in the CSV."""
        results = [_formula_result(formula)]
        filepath = CSVReporter().generate(results, output_dir=str(tmp_path))

        with open(filepath, encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                ioc_val = row["ioc_value"]
                assert ioc_val != formula, (
                    f"Raw formula '{formula}' written unprotected to CSV"
                )
                assert ioc_val.startswith("'"), (
                    f"Expected apostrophe prefix, got: {ioc_val!r}"
                )

    def test_numeric_values_pass_through(self, tmp_path: Path):
        """Numeric abuse scores are written as-is (not wrapped in quotes by our code)."""
        r = EnrichmentResult(ioc=IOC(value="8.8.8.8", ioc_type=IOCType.IP))
        r.abuse_score = 0
        r.sources["test"] = {}
        r.set_verdict()
        filepath = CSVReporter().generate([r], output_dir=str(tmp_path))

        with open(filepath, encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                # abuse_score should be "0", not "'0"
                assert not row["abuse_score"].startswith("'"), (
                    "Numeric zero was incorrectly apostrophe-prefixed"
                )


# ── SQLite cache round-trip integration ───────────────────────────────────────

class TestSQLiteCacheIntegration:

    def test_ip_result_round_trip(self, tmp_path: Path):
        store = InvestigationStore(str(tmp_path / "test.db"))
        r = EnrichmentResult(ioc=IOC(value="1.2.3.4", ioc_type=IOCType.IP))
        r.abuse_score = 85
        r.country = "US"
        r.organization = "Test ISP"
        r.sources["abuseipdb"] = {"abuse_score": 85}
        r.set_verdict()

        store.cache_result(r, expires_at=9_999_999_999)
        cached = store.get_cached(r.ioc, now=1)

        assert cached is not None
        assert cached.ioc.value == "1.2.3.4"
        assert cached.ioc.ioc_type == IOCType.IP
        assert cached.verdict == "Malicious"
        assert cached.abuse_score == 85
        assert cached.country == "US"
        assert cached.organization == "Test ISP"

    def test_cve_result_round_trip(self, tmp_path: Path):
        store = InvestigationStore(str(tmp_path / "test.db"))
        r = _cve_result()
        store.cache_result(r, expires_at=9_999_999_999)
        cached = store.get_cached(r.ioc, now=1)

        assert cached is not None
        assert cached.ioc.value == "CVE-2021-44228"
        assert cached.ioc.ioc_type == IOCType.CVE
        assert cached.cvss_score == 10.0
        assert cached.severity == "CRITICAL"
        assert cached.verdict == "Critical"
        assert cached.cve_description == "Log4Shell remote code execution"

    def test_expired_cache_returns_none(self, tmp_path: Path):
        store = InvestigationStore(str(tmp_path / "test.db"))
        r = _cve_result()
        store.cache_result(r, expires_at=1)   # expired in the past
        cached = store.get_cached(r.ioc, now=9_999_999_999)
        assert cached is None

    def test_cache_upsert_overwrites(self, tmp_path: Path):
        """Writing the same IOC twice keeps only the latest payload."""
        store = InvestigationStore(str(tmp_path / "test.db"))
        ioc = IOC(value="CVE-2021-44228", ioc_type=IOCType.CVE)

        r1 = EnrichmentResult(ioc=ioc)
        r1.cvss_score = 9.0
        r1.sources["nvd"] = {}
        r1.set_verdict()
        store.cache_result(r1, expires_at=9_999_999_999)

        r2 = EnrichmentResult(ioc=ioc)
        r2.cvss_score = 10.0          # updated score
        r2.sources["nvd"] = {}
        r2.set_verdict()
        store.cache_result(r2, expires_at=9_999_999_999)

        cached = store.get_cached(ioc, now=1)
        assert cached is not None
        assert cached.cvss_score == 10.0, "Cache upsert did not overwrite old entry"

    def test_record_investigation_persists(self, tmp_path: Path):
        """record_investigation stores results retrievable via SQLite directly."""
        import sqlite3
        store = InvestigationStore(str(tmp_path / "test.db"))
        results = [_cve_result(), _cve_result("CVE-2023-1234")]
        inv_id = store.record_investigation(results)

        conn = sqlite3.connect(str(tmp_path / "test.db"))
        row = conn.execute(
            "SELECT total_iocs FROM investigations WHERE id = ?", (inv_id,)
        ).fetchone()
        conn.close()

        assert row is not None
        assert row[0] == 2
